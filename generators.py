from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,ListFlowable,ListItem
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

def create_pdf(n,p):
    s=getSampleStyleSheet(); title=ParagraphStyle("T",parent=s["Title"],alignment=TA_CENTER)
    d=SimpleDocTemplate(str(p),pagesize=A4); a=[Paragraph(n.get("title","AI Notes"),title),Spacer(1,12)]
    a += [Paragraph("<b>Summary</b>",s["Heading2"]),Paragraph(n.get("summary",""),s["BodyText"])]
    a += [Spacer(1,8),Paragraph("<b>Key Points</b>",s["Heading2"])]
    a += [ListFlowable([ListItem(Paragraph(str(x),s["BodyText"])) for x in n.get("key_points",[])],bulletType="bullet")]
    a += [Paragraph("<b>Detailed Notes</b>",s["Heading2"])]
    for x in n.get("detailed_notes",[]): a += [Paragraph(str(x.get("heading","")),s["Heading3"]),Paragraph(str(x.get("content","")),s["BodyText"])]
    a += [Paragraph("<b>Important Questions</b>",s["Heading2"])]
    for q in n.get("important_questions",[]): a.append(Paragraph("• "+str(q),s["BodyText"]))
    a += [Paragraph("<b>MCQs</b>",s["Heading2"])]
    for i,x in enumerate(n.get("mcqs",[]),1):
        a.append(Paragraph(f"{i}. {x.get('question','')}",s["BodyText"]))
        for o in x.get("options",[]): a.append(Paragraph("   "+str(o),s["BodyText"]))
        a.append(Paragraph("<b>Answer: </b>"+str(x.get("answer","")),s["BodyText"]))
    d.build(a)

def create_docx(n,p):
    d=Document(); d.add_heading(n.get("title","AI Notes"),0)
    d.add_heading("Summary",1); d.add_paragraph(n.get("summary",""))
    d.add_heading("Key Points",1)
    for x in n.get("key_points",[]): d.add_paragraph(str(x),style="List Bullet")
    d.add_heading("Detailed Notes",1)
    for x in n.get("detailed_notes",[]): d.add_heading(str(x.get("heading","")),2); d.add_paragraph(str(x.get("content","")))
    d.add_heading("Important Questions",1)
    for q in n.get("important_questions",[]): d.add_paragraph(str(q),style="List Number")
    d.add_heading("MCQs",1)
    for i,x in enumerate(n.get("mcqs",[]),1):
        d.add_paragraph(f"{i}. {x.get('question','')}")
        for o in x.get("options",[]): d.add_paragraph(str(o),style="List Bullet 2")
        d.add_paragraph("Answer: "+str(x.get("answer","")))
    d.save(p)

def create_xlsx(n,p):
    w=Workbook(); s=w.active; s.title="Notes"; s["A1"]=n.get("title","AI Notes"); s["A1"].font=Font(bold=True,size=16)
    s["A3"]="Summary"; s["A3"].font=Font(bold=True); s["B3"]=n.get("summary",""); r=5
    s.cell(r,1,"Key Points").font=Font(bold=True); r+=1
    for x in n.get("key_points",[]): s.cell(r,1,x); r+=1
    r+=1; s.cell(r,1,"Detailed Notes").font=Font(bold=True); r+=1; s.append(["Heading","Content"])
    for x in n.get("detailed_notes",[]): r+=1; s.cell(r,1,x.get("heading","")); s.cell(r,2,x.get("content",""))
    q=w.create_sheet("Questions & MCQs"); q.append(["Type","Question","Options","Answer"])
    for x in n.get("important_questions",[]): q.append(["Important Question",x,"",""])
    for x in n.get("mcqs",[]): q.append(["MCQ",x.get("question","")," | ".join(x.get("options",[])),x.get("answer","")])
    w.save(p)
